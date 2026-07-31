from __future__ import annotations

import base64
import csv

import pytest

from app.api import WorktimeApi
from common import database, export


@pytest.fixture(autouse=True)
def isolate_import_logs(tmp_path, monkeypatch):
    monkeypatch.setattr(export, "IMPORT_LOG_DIR", tmp_path / "import_logs")


def make_conn(tmp_path, name="database.db"):
    db_path = tmp_path / name
    database.init_db(db_path)
    conn = database.connect(db_path)
    database.set_settings(
        conn,
        {
            "weekly_target_hours": "40",
            "workdays_per_week": "5",
            "workday_weekdays": "0,1,2,3,4",
            "bundesland": "",
            "daily_break_minutes": "0",
            "tracking_start_date": "2026-07-01",
            "initial_flextime_minutes": "0",
        },
    )
    return conn


def read_csv_rows(path):
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle, delimiter=";"))


def test_detailed_export_contains_segments_absences_notes_and_summary(tmp_path):
    conn = make_conn(tmp_path)
    database.add_segment(conn, "2026-07-06", "WORK", "08:00:00", "12:00:00", "OFFICE", "MANUAL")
    database.add_segment(conn, "2026-07-06", "BREAK", "12:00:00", "12:30:00", source="MANUAL")
    database.upsert_day_type(conn, "2026-07-06", "DIENSTREISE", note="Kunde")
    database.replace_note(conn, "2026-07-06", "Wichtiger Termin")

    path = export.export_period(conn, "2026-07-06", "2026-07-06", "csv", output_dir=tmp_path)
    rows = read_csv_rows(path)
    row_types = [row["Datensatz"] for row in rows]

    assert row_types == ["SEGMENT", "SEGMENT", "ABWESENHEIT", "NOTIZ", "TAGES_SUMME"]
    assert rows[0]["Beginn"] == "08:00"
    assert rows[0]["Ende"] == "12:00"
    assert rows[0]["Dauer_Minuten"] == "240"
    assert rows[2]["Kategorie"] == "DIENSTREISE"
    assert rows[3]["Notiz"] == "Wichtiger Termin"
    assert rows[4]["Soll_Minuten"] == "480"


def test_csv_import_recreates_editable_local_data(tmp_path):
    source = make_conn(tmp_path, "source.db")
    database.add_segment(source, "2026-07-06", "WORK", "08:00:00", "12:00:00", "HOME", "MANUAL")
    database.add_segment(source, "2026-07-06", "WORK", "13:00:00", "17:00:00", "HOME", "MANUAL")
    database.upsert_day_type(source, "2026-07-07", "GLEITZEITTAG", note="Ausgleich")
    database.replace_note(source, "2026-07-06", "Import-Test")
    export_path = export.export_period(source, "2026-07-06", "2026-07-07", "csv", output_dir=tmp_path)

    target = make_conn(tmp_path, "target.db")
    result = export.import_file(target, export_path)

    assert result["segments"] == 2
    assert result["day_types"] == 1
    assert result["notes"] == 1
    assert len(database.get_segments_for_date(target, "2026-07-06")) == 2
    assert database.get_day_types_for_date(target, "2026-07-07")[0]["type"] == "GLEITZEITTAG"
    assert database.get_note_for_date(target, "2026-07-06") == "Import-Test"
    assert database.get_day_summary(target, "2026-07-06")["actual_minutes"] == 480


def test_xlsx_import_recreates_editable_local_data(tmp_path):
    source = make_conn(tmp_path, "xlsx-source.db")
    database.add_segment(source, "2026-07-06", "WORK", "08:15:00", "16:45:00", "OFFICE", "MANUAL")
    database.replace_note(source, "2026-07-06", "Excel-Test")
    export_path = export.export_period(source, "2026-07-06", "2026-07-06", "xlsx", output_dir=tmp_path)

    target = make_conn(tmp_path, "xlsx-target.db")
    result = export.import_file(target, export_path)

    assert result["segments"] == 1
    assert result["notes"] == 1
    assert database.get_segments_for_date(target, "2026-07-06")[0]["start_time"] == "08:15:00"
    assert database.get_note_for_date(target, "2026-07-06") == "Excel-Test"


def test_xlsx_export_has_readable_overview_and_import_sheet(tmp_path):
    conn = make_conn(tmp_path, "overview-source.db")
    database.add_segment(conn, "2026-07-06", "WORK", "08:15:00", "16:45:00", "OFFICE", "MANUAL")
    database.upsert_day_type(conn, "2026-07-06", "GLEITZEITTAG", note="Ausgleich")
    database.replace_note(conn, "2026-07-06", "Excel-Test")
    export_path = export.export_period(conn, "2026-07-06", "2026-07-06", "xlsx", output_dir=tmp_path)

    from openpyxl import load_workbook

    workbook = load_workbook(export_path, read_only=True, data_only=True)

    assert workbook.sheetnames[:5] == ["Übersicht", "Segmente", "Abwesenheiten", "Notizen", "Importdaten"]
    assert [cell.value for cell in next(workbook["Übersicht"].iter_rows(max_row=1))][:4] == ["Datum", "Wochentag", "Tagesart", "Beginn"]
    assert [cell.value for cell in next(workbook["Segmente"].iter_rows(max_row=1))][:4] == ["Datum", "Wochentag", "Typ", "Beginn"]
    assert next(workbook["Segmente"].iter_rows(min_row=2, max_row=2, values_only=True))[2] == "Arbeit"
    assert [cell.value for cell in next(workbook["Importdaten"].iter_rows(max_row=1))][:4] == ["Datensatz", "Datum", "Wochentag", "Kategorie"]
    assert [row[0] for row in workbook["Importdaten"].iter_rows(min_row=2, values_only=True)] == ["SEGMENT", "ABWESENHEIT", "NOTIZ"]


def test_xlsx_import_can_read_readable_sheets(tmp_path):
    source = make_conn(tmp_path, "readable-source.db")
    database.add_segment(source, "2026-07-06", "WORK", "08:15:00", "16:45:00", "HOME", "MANUAL")
    database.upsert_day_type(source, "2026-07-07", "GLEITZEITTAG", note="Ausgleich")
    database.replace_note(source, "2026-07-06", "Lesbare Tabellen")
    export_path = export.export_period(source, "2026-07-06", "2026-07-07", "xlsx", output_dir=tmp_path)

    target = make_conn(tmp_path, "readable-target.db")
    result = export.import_file(target, export_path)

    assert result["segments"] == 1
    assert result["day_types"] == 1
    assert result["notes"] == 1
    assert database.get_segments_for_date(target, "2026-07-06")[0]["location"] == "HOME"
    assert database.get_day_types_for_date(target, "2026-07-07")[0]["type"] == "GLEITZEITTAG"


def test_xlsx_import_creates_work_segment_from_overview_when_no_detail_segments_exist(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Übersicht"
    sheet.append(export.SUMMARY_HEADERS)
    sheet.append(["2026-07-06", "Montag", "Arbeitstag", "08:00", "16:00", "", "", "", "", "Büro", "Nachtrag"])
    path = tmp_path / "overview-only.xlsx"
    workbook.save(path)

    target = make_conn(tmp_path, "overview-only-target.db")
    result = export.import_file(target, path)

    assert result["overview_segments"] == 1
    assert result["overview_notes"] == 1
    assert result["warnings"] == 0
    assert result["log_name"].endswith(".html")
    assert (tmp_path / "import_logs" / result["log_name"]).exists()
    segment = database.get_segments_for_date(target, "2026-07-06")[0]
    assert segment["type"] == "WORK"
    assert segment["start_time"] == "08:00:00"
    assert segment["end_time"] == "16:00:00"
    assert segment["location"] == "OFFICE"
    assert database.get_note_for_date(target, "2026-07-06") == "Nachtrag"


def test_xlsx_import_uses_detail_segments_before_overview_and_logs_mismatch(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Übersicht"
    overview.append(export.SUMMARY_HEADERS)
    overview.append(["2026-07-06", "Montag", "Arbeitstag", "08:00", "16:00", "", "", "", "", "Büro", ""])
    segments = workbook.create_sheet("Segmente")
    segments.append(export.SEGMENT_SHEET_HEADERS)
    segments.append(["2026-07-06", "Montag", "Arbeit", "10:00", "16:00", "", "Büro", "MANUAL"])
    path = tmp_path / "overview-detail-mismatch.xlsx"
    workbook.save(path)

    target = make_conn(tmp_path, "overview-detail-target.db")
    result = export.import_file(target, path)

    assert result["segments"] == 1
    assert result["overview_segments"] == 0
    assert result["warnings"] >= 1
    segment = database.get_segments_for_date(target, "2026-07-06")[0]
    assert segment["start_time"] == "10:00:00"
    assert segment["end_time"] == "16:00:00"
    log_content = (tmp_path / "import_logs" / result["log_name"]).read_text(encoding="utf-8")
    assert "Übersicht passt nicht zu Segmenten" in log_content
    assert "const IMPORT_GROUPS" in log_content
    assert '"status_label": "Warnung"' in log_content
    assert "Vorher" in log_content
    assert "Nachher" in log_content


def test_import_replaces_only_dates_contained_in_file(tmp_path):
    source = make_conn(tmp_path, "replace-source.db")
    database.add_segment(source, "2026-07-06", "WORK", "09:00:00", "17:00:00", "HOME", "MANUAL")
    export_path = export.export_period(source, "2026-07-06", "2026-07-06", "csv", output_dir=tmp_path)

    target = make_conn(tmp_path, "replace-target.db")
    database.add_segment(target, "2026-07-05", "WORK", "08:00:00", "12:00:00", "OFFICE", "MANUAL")
    database.add_segment(target, "2026-07-06", "WORK", "06:00:00", "07:00:00", "OFFICE", "MANUAL")

    export.import_file(target, export_path)

    assert len(database.get_segments_for_date(target, "2026-07-05")) == 1
    imported_day = database.get_segments_for_date(target, "2026-07-06")
    assert len(imported_day) == 1
    assert imported_day[0]["start_time"] == "09:00:00"
    assert imported_day[0]["location"] == "HOME"


def test_api_imports_uploaded_base64_file_without_native_dialog(tmp_path):
    source = make_conn(tmp_path, "upload-source.db")
    database.add_segment(source, "2026-07-06", "WORK", "09:00:00", "17:00:00", "HOME", "MANUAL")
    export_path = export.export_period(source, "2026-07-06", "2026-07-06", "csv", output_dir=tmp_path)

    api = WorktimeApi(tmp_path / "upload-target.db")
    payload = base64.b64encode(export_path.read_bytes()).decode("ascii")
    result = api.import_uploaded_file(export_path.name, payload)

    assert result["segments"] == 1
    with database.connect(api.db_path) as conn:
        assert database.get_segments_for_date(conn, "2026-07-06")[0]["location"] == "HOME"


def test_sap_sdata_preview_pairs_events_and_imports_selected_days(tmp_path):
    conn = make_conn(tmp_path, "sap.db")
    path = tmp_path / "sap.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["SDATA"], delimiter=";")
        writer.writeheader()
        for event_type, log_time in (
            ("P10", "100000"),
            ("P20", "120000"),
            ("P10", "130000"),
            ("P20", "170000"),
        ):
            writer.writerow({"SDATA": sap_sdata(event_type, "20260101", log_time)})

    preview = export.preview_sdata_file(conn, path)

    day = preview["days"][0]
    assert day["date"] == "2026-01-01"
    assert [segment["type"] for segment in day["imported_segments"]] == ["WORK", "BREAK", "WORK"]
    assert [segment["location"] for segment in day["imported_segments"]] == ["OFFICE", "", "OFFICE"]
    assert [(segment["start_time"], segment["end_time"]) for segment in day["imported_segments"]] == [
        ("10:00:00", "12:00:00"),
        ("12:00:00", "13:00:00"),
        ("13:00:00", "17:00:00"),
    ]

    result = export.import_sdata_preview(conn, preview["days"], ["2026-01-01"], source_name="SAP-Test")
    segments = database.get_segments_for_date(conn, "2026-01-01")

    assert result["segments"] == 3
    assert [segment["type"] for segment in segments] == ["WORK", "BREAK", "WORK"]
    assert [segment["location"] for segment in segments] == ["OFFICE", None, "OFFICE"]


def sap_sdata(event_type: str, log_date: str, log_time: str) -> str:
    return f"PP2CLNT080{event_type}0131{log_date}{log_time}{log_date}{log_time}        00002634"
